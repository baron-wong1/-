import hashlib
import json
import shutil
import subprocess
import sys
from decimal import Decimal
from pathlib import Path

import fitz
import pytest
from openpyxl import load_workbook
from PIL import Image

from invoice_assistant.parser import parse
from invoice_assistant.models import money, reconcile
from invoice_assistant.importer import import_files, preview, recognize, digest, discover, soffice_path
from invoice_assistant.exporter import export_files, prepare


def imported(paths, tmp_path, records=None):
    return import_files({'paths':list(map(str, paths)), 'cache':str(tmp_path / 'cache'), 'records':records or []}, lambda *a, **k:None)['records']


def reviewed(records):
    return reconcile([{**r, 'reviewed':True} for r in records])


def request(rows, tmp_path):
    return dict(records=rows, person='测试人员', month='2026-09', destination=str(tmp_path / 'out'))


def test_parser_preserves_identifier_and_tax_total():
    result=parse('发票号码：00123456789012345678\n开票日期：2026年9月3日\n金额 100.00\n税额 23.45\n价税合计（小写）￥123.45\n酒店')
    assert result['number']=='00123456789012345678'
    assert result['amount']=='123.45'
    assert result['date']=='2026-09-03'
    assert result['category']=='住宿'


def test_multi_ticket_ambiguity_never_reuses_first_amount():
    result=parse('发票号码:00123456789012345678\n开票日期:2026-09-03\n票价:100.00\n发票号码:00123456789012345679\n票价:50.00')
    assert result['number']=='' and result['amount'] is None and result['date'] is None
    assert any('多个发票号码' in w for w in result['warnings'])


def test_missing_and_invalid_fields_are_not_zero():
    result=parse('这是一份普通文档')
    assert result['amount'] is None and result['date'] is None
    assert money(None) is None
    for bad in ['NaN','Infinity','-1','1.234','1e99']:
        with pytest.raises(ValueError):money(bad)
    assert sum([money('0.10'),money('0.20')])==Decimal('0.30')


def test_file_hash_identity_and_same_name_different_content(pdf_factory,tmp_path):
    a=pdf_factory()
    folder=tmp_path/'second';folder.mkdir()
    b=pdf_factory(name='second/invoice.pdf',number='00123456789012345679',amount='88.00')
    rows=imported([a,b],tmp_path)
    assert len(rows)==2 and len({r['source_hash'] for r in rows})==2
    assert len(imported([a],tmp_path,rows))==2


def test_errors_isolated_from_good_files(pdf_factory,tmp_path):
    good=pdf_factory()
    broken=tmp_path/'broken.pdf';broken.write_bytes(b'not pdf')
    rows=imported([broken,good],tmp_path)
    assert len(rows)==2
    assert len([r for r in rows if r['status']=='error'])==1
    assert len([r for r in rows if r['status']=='ready'])==1


def test_encrypted_pdf_visible_error(tmp_path):
    doc=fitz.open();doc.new_page();path=tmp_path/'locked.pdf'
    doc.save(path,encryption=fitz.PDF_ENCRYPT_AES_256,user_pw='secret',owner_pw='owner');doc.close()
    rows=imported([path],tmp_path)
    assert rows[0]['status']=='error' and '加密' in rows[0]['error']


def test_recursive_discovery_excludes_outputs_and_symlinks(pdf_factory,tmp_path):
    good=pdf_factory()
    output=tmp_path/'整理结果';output.mkdir();shutil.copy(good,output/'copy.pdf')
    try:
        (tmp_path/'loop').symlink_to(tmp_path,target_is_directory=True)
    except OSError:
        pass  # Some Windows runners disallow creating symlinks.
    files,errors=discover([tmp_path],[])
    assert files==[good] and not errors


def test_duplicate_and_conflict_can_be_resolved(pdf_factory,tmp_path):
    a=pdf_factory()
    b=pdf_factory(name='second.pdf',text='电子发票\n发票号码:00123456789012345678\n开票日期:2026年09月03日\n住宿服务\n价税合计(小写):￥123.45\n副本')
    rows=reviewed(imported([a,b],tmp_path))
    assert [r['status'] for r in rows]==['ready','duplicate']
    rows[1]['amount']='124.00'
    conflicts=reconcile(rows)
    assert all(r['status']=='conflict' for r in conflicts)
    with pytest.raises(ValueError):prepare(conflicts)
    resolved=reconcile([{**r,'resolution':'keep'} for r in conflicts])
    assert all(r['status']=='ready' for r in resolved)
    assert reconcile([resolved[1]])[0]['status']=='ready'


def test_invalid_ranges_and_crop_block_export(pdf_factory,tmp_path):
    rows=imported([pdf_factory()],tmp_path)
    rows[0]['page_end']=2
    assert reconcile(rows)[0]['status']=='pending'
    rows[0]['page_end']=1;rows[0]['crop']=[0,0,1.1,1]
    assert reconcile(rows)[0]['status']=='pending'


def test_complete_export_amount_identifiers_layout_bookmarks(pdf_factory,tmp_path):
    a=pdf_factory(day='2026年09月05日',amount='0.10')
    b=pdf_factory(name='b.pdf',number='00123456789012345679',day='2026年09月03日',amount='0.20')
    rows=imported([a,b],tmp_path)
    before={p:digest(p) for p in [a,b]}
    result=export_files(request(rows,tmp_path),lambda *a,**k:None)
    assert result['count']==2 and result['amount']=='0.30' and result['pages']==1
    combined=next(Path(p) for p in result['files'] if '全部发票' in p)
    with fitz.open(combined) as doc:
        assert doc.page_count==1
        assert doc.get_toc()==[[1,'住宿',1]]
        text=doc[0].get_text()
        assert text.index(rows[1]['number'])<text.index(rows[0]['number'])
        assert '1 / 1' in text
        assert doc[0].rect.width==pytest.approx(595.2756,abs=.001)
    workbook=load_workbook(next(p for p in result['files'] if p.endswith('.xlsx')))
    sheet=workbook.active
    assert sheet['F3'].value=='00123456789012345679' and sheet['F3'].data_type=='s'
    assert sheet['F4'].value=='00123456789012345678' and sheet['F4'].number_format=='@'
    assert sheet['D5'].value==0.3
    assert [sheet.cell(2,c).value for c in range(1,7)]==['报销日期','报销人','发票内容','发票金额','开票日期','发票号码']
    assert all(digest(p)==h for p,h in before.items())
    # A second export never replaces previous output.
    second=export_files(request(rows,tmp_path),lambda *a,**k:None)
    assert second['directory']!=result['directory'] and combined.exists()


def test_attachment_half_page_excluded_from_excel(pdf_factory,tmp_path):
    invoice=pdf_factory()
    image=tmp_path/'payment.png';Image.new('RGB',(400,1200),'#4a8fc4').save(image)
    attachment=imported([image],tmp_path)[0]
    attachment.update(kind='attachment',category='住宿',reviewed=True,warnings=[])
    rows=reviewed([*imported([invoice],tmp_path),attachment])
    result=export_files(request(rows,tmp_path),lambda *a,**k:None)
    assert result['count']==1 and result['amount']=='123.45'
    with fitz.open(next(p for p in result['files'] if '全部发票' in p)) as doc:
        images=doc[0].get_image_info()
        assert images and images[0]['bbox'][3]-images[0]['bbox'][1] <= doc[0].rect.height/2


def test_long_invoice_single_page(pdf_factory,tmp_path):
    rows=imported([pdf_factory(size=(300,650)),pdf_factory(name='b.pdf',number='00123456789012345679')],tmp_path)
    result=export_files(request(rows,tmp_path),lambda *a,**k:None)
    assert result['pages']==2


def test_multpage_ticket_counted_once_and_overlaps_rejected(pdf_factory,tmp_path):
    rows=reviewed(imported([pdf_factory(pages=2)],tmp_path))
    assert rows[0]['warnings']
    merged={**rows[0],'page_end':2}
    result=export_files(request([merged],tmp_path),lambda *a,**k:None)
    assert result['count']==1
    second={**merged,'id':'new','number':'00123456789012345679'}
    with pytest.raises(ValueError,match='重叠'):prepare([merged,second])
    merged['page_end']=1;merged['crop']=[0,0,1,.5];second['page_end']=1;second['crop']=[0,.5,1,1]
    assert len(prepare([merged,second])[0])==2


def test_changed_original_blocks_preview_and_export(pdf_factory,tmp_path):
    path=pdf_factory();rows=imported([path],tmp_path)
    path.write_bytes(b'changed')
    with pytest.raises(RuntimeError,match='原文件已变更'):preview({'record':rows[0]})
    with pytest.raises(RuntimeError,match='原文件已变更'):export_files(request(rows,tmp_path),lambda *a,**k:None)
    assert not list((tmp_path/'out').glob('发票整理_*'))


def test_preview_and_crop_recognition(pdf_factory,tmp_path):
    rows=imported([pdf_factory()],tmp_path)
    p=preview({'record':rows[0]})
    assert p['image'].startswith('data:image/png;base64,') and p['width']>0
    r=recognize({'record':{**rows[0],'crop':[0,0,1,1]}})['record']
    assert r['number']=='00123456789012345678' and r['amount']=='123.45'


def test_scanned_pdf_and_image_use_offline_ocr(pdf_factory,tmp_path):
    original=pdf_factory(size=(800,350))
    with fitz.open(original) as doc:
        pix=doc[0].get_pixmap(matrix=fitz.Matrix(2,2))
        png=tmp_path/'scan.png';pix.save(png)
        scanned=fitz.open();page=scanned.new_page(width=800,height=350);page.insert_image(page.rect,stream=pix.tobytes('png'))
        scanpdf=tmp_path/'scan.pdf';scanned.save(scanpdf);scanned.close()
    rows=imported([png,scanpdf],tmp_path)
    assert len(rows)==2
    for row in rows:
        assert row['method']=='ocr'
        assert row['number']=='00123456789012345678'
        assert row['amount']=='123.45'
        assert row['status'] in ('pending','conflict') and row['warnings']


def test_docx_and_doc_conversion(docx_factory,tmp_path):
    if not soffice_path():pytest.skip('LibreOffice unavailable')
    path=docx_factory()
    rows=reviewed(imported([path],tmp_path))
    assert rows[0]['number']=='00123456789012345678' and rows[0]['amount']=='123.45'
    output=tmp_path/'legacy';output.mkdir()
    subprocess.run([str(soffice_path()),'-env:UserInstallation='+(tmp_path/'doc-profile').as_uri(),'--headless',
                    '--convert-to','doc:MS Word 97','--outdir',str(output),str(path)],check=True,timeout=120,capture_output=True)
    legacy=output/'invoice.doc'
    assert legacy.exists()
    converted=reviewed(imported([legacy],tmp_path))
    assert converted[0]['number']=='00123456789012345678' and converted[0]['amount']=='123.45'


def test_excel_user_strings_not_formulas(pdf_factory,tmp_path):
    rows=imported([pdf_factory()],tmp_path)
    rows[0]['category']='=1+1'
    r=request(rows,tmp_path);r['person']='=2+2'
    result=export_files(r,lambda *a,**k:None)
    wb=load_workbook(next(p for p in result['files'] if p.endswith('.xlsx')))
    assert wb.active['B3'].data_type=='s' and wb.active['C3'].data_type=='s'
    assert wb.active['B3'].value=='=2+2'


def test_worker_protocol(pdf_factory,tmp_path):
    worker=Path(__file__).resolve().parents[1]/'core/worker.py'
    proc=subprocess.run([sys.executable,str(worker)],input=json.dumps({'op':'import','paths':[str(pdf_factory())],
                        'cache':str(tmp_path/'cache')})+'\n',capture_output=True,text=True,encoding='utf-8',check=True)
    events=[json.loads(line) for line in proc.stdout.splitlines()]
    assert events[0]['event']=='started' and events[-1]['event']=='result'
    assert events[-1]['result']['records'][0]['amount']=='123.45'


def test_fullwidth_ocr_decimals_and_invalid_precision():
    parsed=parse('发票号码：００１２３４５６７８９０１２３４５６７８\n开票日期：２０２６年９月３日\n价税合计（小写）：￥１２３．４５')
    assert parsed['amount']=='123.45' and parsed['number']=='00123456789012345678'
    assert parse('价税合计(小写):￥123.456')['amount'] is None


def test_multiple_invoice_page_cannot_be_confirmed_as_single_whole_page(pdf_factory,tmp_path):
    row=imported([pdf_factory()],tmp_path)[0]
    row.update(requires_split=True,reviewed=True)
    assert reconcile([row])[0]['status']=='pending'


def test_category_conflict_is_not_silently_deduplicated(pdf_factory,tmp_path):
    row=imported([pdf_factory()],tmp_path)[0]
    other={**row,'id':'second','category':'车票'}
    assert all(r['status']=='conflict' for r in reconcile([row,other]))


def test_rotated_pdf_normalized_without_changing_original(pdf_factory,tmp_path):
    path=pdf_factory()
    with fitz.open(path) as doc:
        doc[0].set_rotation(90)
        doc.saveIncr()
        original_size=doc[0].rect.width,doc[0].rect.height
    before=digest(path)
    rows=reviewed(imported([path],tmp_path))
    assert not rows[0].get('error')
    with fitz.open(rows[0]['render_path']) as normalized:
        assert normalized[0].rotation==0
        assert (normalized[0].rect.width,normalized[0].rect.height)==original_size
    assert rows[0]['number']=='00123456789012345678'
    export_files(request(rows,tmp_path),lambda *a,**k:None)
    assert digest(path)==before
