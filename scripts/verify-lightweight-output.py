"""Independent read-back with libraries used only in CI, never shipped."""
from pathlib import Path
from decimal import Decimal
import sys
import fitz
from openpyxl import load_workbook

root = Path(sys.argv[1])
output = next((root / 'exports').iterdir())
assert len(list(output.glob('*.pdf'))) == 3
with fitz.open(next(output.glob('*合并打印.pdf'))) as doc:
    assert len(doc) == 2 and len(doc.get_toc()) == 2
    for i, page in enumerate(doc):
        assert abs(page.rect.width - 595.2756) < 0.1
        assert f'{i + 1} / 2' in page.get_text()
        page.get_pixmap(matrix=fitz.Matrix(1.4, 1.4)).save(root / f'printed-page-{i + 1}.png')
    assert '00123456789012345678' in doc[0].get_text()
    assert '00888888888888888888' in doc[1].get_text()
wb = load_workbook(next(output.glob('*.xlsx')))
s = wb.active
assert [s.cell(2, i).value for i in range(1, 7)] == ['报销日期', '报销人', '发票内容', '发票金额', '开票日期', '发票号码']
assert s['F3'].value == '00123456789012345678' and s['F3'].data_type == 's'
assert s['B3'].value == '=SUM(A1:A2)' and s['B3'].data_type == 's'
assert Decimal(str(s['D3'].value)) + Decimal(str(s['D4'].value)) == Decimal('133.55')
assert Decimal(str(s['D5'].value)) == Decimal('133.55')
assert s['E3'].value.strftime('%Y-%m-%d') == '2026-10-03'
assert all(c.data_type != 'f' for row in s for c in row)
print('Independent PDF/Excel read-back passed; 2 invoice rows, 133.55 total, 1 excluded attachment.')
