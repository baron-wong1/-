from __future__ import annotations

import os
import re
import shutil
import tempfile
from collections import OrderedDict
from datetime import datetime
from decimal import Decimal
from pathlib import Path

import fitz
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from .importer import check_source, clip_rect
from .models import CATEGORIES, money, reconcile, safe_name, sort_key, inside

WIDTH, HEIGHT = 595.2756, 841.8898
MARGIN, FOOTER, GAP = 24, 25, 12


def prepare(records):
    rows = reconcile(records)
    blocked = [row for row in rows if row['status'] in ('pending', 'conflict', 'error')]
    if blocked:
        raise ValueError(f'还有 {len(blocked)} 项需要核对或移除，未生成文件')
    active = [row for row in rows if row['status'] == 'ready']
    if not any(row['kind'] == 'invoice' for row in active):
        raise ValueError('请至少导入并确认一张发票')
    # Page ranges may overlap only when the user selected disjoint rectangles.
    for i, row in enumerate(active):
        for other in active[i + 1:]:
            if row['source_hash'] != other['source_hash']:
                continue
            if max(row['page_start'], other['page_start']) > min(row['page_end'], other['page_end']):
                continue
            a, b = row.get('crop') or [0, 0, 1, 1], other.get('crop') or [0, 0, 1, 1]
            if max(a[0], b[0]) < min(a[2], b[2]) - 0.001 and max(a[1], b[1]) < min(a[3], b[3]) - 0.001:
                raise ValueError(f"{row['filename']} 的票据页范围或裁剪区域重叠，请拆分或移除多余记录")
    return active, [r for r in rows if r['status'] == 'duplicate']


def draw_piece(output_page, doc, page_index, row, area):
    source = doc[page_index]
    clip = clip_rect(source, row.get('crop'))
    # show_pdf_page preserves text, QR codes and seals. No automatic content cropping.
    output_page.show_pdf_page(area, doc, page_index, clip=clip, keep_proportion=True)


def make_pdf(rows):
    output = fitz.open()
    pending = []
    area_width = WIDTH - MARGIN * 2
    half_height = (HEIGHT - MARGIN - FOOTER - GAP) / 2

    def flush():
        if not pending:
            return
        page = output.new_page(width=WIDTH, height=HEIGHT)
        for index, (row, doc, pi) in enumerate(pending):
            top = MARGIN + index * (half_height + GAP)
            draw_piece(page, doc, pi, row, fitz.Rect(MARGIN, top, WIDTH - MARGIN, top + half_height))
        for _, doc, _ in pending:
            doc.close()
        pending.clear()

    for row in sorted(rows, key=sort_key):
        for page_index in range(row['page_start'] - 1, row['page_end']):
            doc = fitz.open(row['render_path'])
            clip = clip_rect(doc[page_index], row.get('crop'))
            scaled_height = clip.height * area_width / clip.width
            long_invoice = row['kind'] == 'invoice' and scaled_height > half_height
            if long_invoice:
                flush()
                page = output.new_page(width=WIDTH, height=HEIGHT)
                draw_piece(page, doc, page_index, row, fitz.Rect(MARGIN, MARGIN, WIDTH - MARGIN, HEIGHT - FOOTER))
                doc.close()
            else:
                pending.append((row, doc, page_index))
                if len(pending) == 2:
                    flush()
    flush()
    return output


def add_page_numbers(doc):
    for index, page in enumerate(doc):
        label = f'{index + 1} / {doc.page_count}'
        x = WIDTH - MARGIN - fitz.get_text_length(label, fontsize=8)
        page.insert_text((x, HEIGHT - 12), label, fontsize=8, color=(0.4, 0.4, 0.4))


def text_cell(cell, value):
    # Do not let user-entered strings become spreadsheet formulas.
    cell.value = str(value)
    cell.data_type = 's'


def make_excel(path, rows, person, month, categories):
    invoices = [row for category in categories for row in sorted(rows, key=sort_key)
                if row['category'] == category and row['kind'] == 'invoice']
    wb = Workbook()
    sheet = wb.active
    sheet.title = '电子发票明细'
    sheet.merge_cells('A1:F1')
    text_cell(sheet['A1'], f'{month[:4]}年{int(month[5:])}月电子发票明细表')
    sheet['A1'].font = Font(name='宋体', size=16, bold=True)
    sheet['A1'].alignment = Alignment(horizontal='center', vertical='center')
    sheet.row_dimensions[1].height = 32
    headers = ['报销日期', '报销人', '发票内容', '发票金额', '开票日期', '发票号码']
    sheet.append(headers)
    for row in invoices:
        sheet.append([f'{month[:4]}年{int(month[5:])}月', person, row['category'], money(row['amount']),
                      datetime.strptime(row['date'], '%Y-%m-%d'), row['number']])
        i = sheet.max_row
        for col in (1, 2, 3, 6):
            text_cell(sheet.cell(i, col), sheet.cell(i, col).value)
        sheet.cell(i, 4).number_format = '#,##0.00'
        sheet.cell(i, 5).number_format = 'yyyy-m-d'
        sheet.cell(i, 6).number_format = '@'
    total = sum((money(row['amount']) for row in invoices), Decimal('0.00'))
    last = sheet.max_row
    sheet.append(['合计', None, None, total])
    sheet.merge_cells(start_row=last + 1, start_column=1, end_row=last + 1, end_column=3)
    sheet.cell(last + 1, 4).number_format = '#,##0.00'
    sheet.cell(last + 1, 4).font = Font(name='宋体', bold=True)
    sheet.append([])
    sheet.append(['类别', '发票张数', '发票金额'])
    summary_header = sheet.max_row
    for category in categories:
        items = [row for row in invoices if row['category'] == category]
        if items:
            sheet.append([category, len(items), sum((money(r['amount']) for r in items), Decimal('0.00'))])
            text_cell(sheet.cell(sheet.max_row, 1), category)
            sheet.cell(sheet.max_row, 3).number_format = '#,##0.00'
    border = Border(bottom=Side(style='thin', color='DCE2E9'))
    for row in sheet.iter_rows():
        for cell in row:
            if cell.row != 1:
                cell.font = Font(name='宋体', size=11, bold=cell.row in (2, last + 1, summary_header))
            cell.alignment = Alignment(horizontal='center' if cell.column != 4 else 'right', vertical='center')
            cell.border = border
        sheet.row_dimensions[row[0].row].height = 25 if row[0].row != 1 else 32
    for i in (2, summary_header):
        for cell in sheet[i]:
            cell.fill = PatternFill('solid', fgColor='EDF2F7')
    for col, width in zip('ABCDEF', (18, 14, 22, 18, 18, 32)):
        sheet.column_dimensions[col].width = width
    sheet.freeze_panes = 'A3'
    sheet.auto_filter.ref = f'A2:F{last}'
    sheet.print_area = f'A1:F{sheet.max_row}'
    sheet.print_title_rows = '1:2'
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight = 1, 0
    wb.save(path)
    return len(invoices), total


def export_files(request, emit):
    person, month = request.get('person', '').strip(), request.get('month', '')
    if not person or len(person) > 80:
        raise ValueError('请填写报销人（最多 80 个字符）')
    if not re.fullmatch(r'20\d{2}-(?:0[1-9]|1[0-2])', month):
        raise ValueError('请选择有效报销月份')
    rows, duplicates = prepare(request['records'])
    destination = Path(request['destination']).resolve()
    destination.mkdir(parents=True, exist_ok=True)
    if any(inside(destination, p) for p in request.get('excluded', [])):
        raise ValueError('不能在程序或缓存目录中生成输出')
    checked = set()
    for row in rows:
        key = row['source_hash'], row['render_path']
        if key not in checked:
            check_source(row)
            checked.add(key)
    categories = [category for category in CATEGORIES if any(r['category'] == category for r in rows)]
    categories.extend(sorted({r['category'] for r in rows} - set(categories)))
    prefix = f'{month}_{safe_name(person)}'
    stamp = datetime.now().strftime('%Y%m%d_%H%M%S_%f')
    output_dir = destination / ('发票整理_' + stamp)
    stage = Path(tempfile.mkdtemp(prefix='.invoice-export-', dir=destination))
    emit('staging', directory=str(stage))
    files, summary = [], []
    try:
        combined = fitz.open()
        bookmarks = []
        for i, category in enumerate(categories):
            emit('progress', current=i + 1, total=len(categories) + 1, filename=f'生成 {category} PDF')
            items = [row for row in rows if row['category'] == category]
            doc = make_pdf(items)
            bookmarks.append([1, category, combined.page_count + 1])
            combined.insert_pdf(doc)
            # Numeric prefix prevents custom category names colliding after sanitization.
            name = f'{i + 1:02d}_{safe_name(category)}_打印版.pdf'
            add_page_numbers(doc)
            doc.save(stage / name, garbage=4, deflate=True)
            invoices = [row for row in items if row['kind'] == 'invoice']
            summary.append(dict(category=category, count=len(invoices), pages=doc.page_count,
                                amount=str(sum((money(r['amount']) for r in invoices), Decimal('0.00')))))
            doc.close()
            files.append(name)
        combined.set_toc(bookmarks)
        add_page_numbers(combined)
        combined_name = prefix + '_全部发票打印版.pdf'
        combined.save(stage / combined_name, garbage=4, deflate=True)
        pages = combined.page_count
        combined.close()
        files.append(combined_name)
        emit('progress', current=len(categories) + 1, total=len(categories) + 1, filename='生成 Excel 明细与分类汇总')
        excel_name = prefix + '_电子发票明细表.xlsx'
        count, total = make_excel(stage / excel_name, rows, person, month, categories)
        files.append(excel_name)
        os.rename(stage, output_dir)
        return dict(directory=str(output_dir), files=[str(output_dir / name) for name in files],
                    count=count, amount=str(total), pages=pages, categories=summary, duplicates=len(duplicates))
    finally:
        if stage.exists():
            shutil.rmtree(stage)
